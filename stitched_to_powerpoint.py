from pathlib import Path
from difflib import SequenceMatcher
from io import BytesIO
import re

import numpy as np
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt


SOURCE = Path(
    r"R:\Basic_Sciences\Phys\SenzaiLab\Aparna\Data_by_mouse"
    r"\m002\histology\collected_stitched"
)
CHANNELS = ("red", "blue", "overlay")

FUZZY_THRESHOLD = 0.75

KNOWN_SPELLINGS = {
    "stitched",
    "stiched",
    "sticthed",
    "stithced",
    "stithed",
    "stitche",
    "stitch",
}


def match_reason(path, channel):
    """Match stitched-channel, allowing spelling/separator variations."""
    stem = path.stem.lower()

    candidates = re.findall(
        rf"([a-z]+)[\s_-]*{re.escape(channel)}(?=$|[^a-z])",
        stem,
    )

    for candidate in candidates:
        if candidate == "stitched":
            return "exact"

        if candidate in KNOWN_SPELLINGS:
            return f"spelling variation: {candidate}"

        similarity = SequenceMatcher(
            None, candidate, "stitched"
        ).ratio()

        if similarity >= FUZZY_THRESHOLD:
            return f"fuzzy match: {candidate} ({similarity:.0%})"

    return None

def natural_sort_key(path):
    """Sort numbers naturally: section2 comes before section10."""
    relative_path = str(path.relative_to(SOURCE)).lower()
    return [
        int(part) if part.isdigit() else part
        for part in re.split(r"(\d+)", relative_path)
    ]


def make_slide_title(path):
    """Use the four folders immediately above the TIFF."""
    return " / ".join(path.parent.parts[-4:])


def tiff_to_png(path):
    """Convert the first TIFF page to a PowerPoint-compatible PNG."""
    with Image.open(path) as original:
        page_count = getattr(original, "n_frames", 1)
        original.seek(0)
        image = ImageOps.exif_transpose(original)

        # Scale high-bit-depth grayscale TIFFs for display, avoiding
        # clipping when converting them directly to 8-bit images.
        if image.mode.startswith("I") or image.mode == "F":
            pixels = np.asarray(image, dtype=np.float64)
            finite = np.isfinite(pixels)

            if finite.any():
                low = pixels[finite].min()
                high = pixels[finite].max()

                if high > low:
                    pixels = (pixels - low) / (high - low) * 255
                else:
                    pixels = np.zeros_like(pixels)
            else:
                pixels = np.zeros_like(pixels)

            pixels = np.nan_to_num(
                pixels, nan=0, posinf=255, neginf=0
            )
            image = Image.fromarray(
                np.clip(pixels, 0, 255).astype(np.uint8)
            ).convert("RGB")
        else:
            image = image.convert("RGB")

        width, height = image.size
        stream = BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)

    return stream, width, height, page_count


def add_slide(presentation, path):
    stream, image_width, image_height, page_count = tiff_to_png(path)

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]  # Blank slide
    )

    title_box = slide.shapes.add_textbox(
        Inches(0.35), Inches(0.15), Inches(12.63), Inches(0.85)
    )
    text_frame = title_box.text_frame
    text_frame.word_wrap = True

    paragraph = text_frame.paragraphs[0]
    paragraph.text = make_slide_title(path)
    paragraph.font.size = Pt(20)
    paragraph.font.bold = True

    # Fit the entire image below the title, preserving its proportions.
    box_left = Inches(0.35)
    box_top = Inches(1.15)
    box_width = Inches(12.63)
    box_height = Inches(6.0)

    scale = min(
        box_width / image_width,
        box_height / image_height,
    )
    picture_width = int(image_width * scale)
    picture_height = int(image_height * scale)

    left = box_left + (box_width - picture_width) // 2
    top = box_top + (box_height - picture_height) // 2

    slide.shapes.add_picture(
        stream,
        left,
        top,
        width=picture_width,
        height=picture_height,
    )

    if page_count > 1:
        print(f"  Note: {page_count} TIFF pages; using the first.")


def main():
    if not SOURCE.is_dir():
        raise FileNotFoundError(f"Source folder not found: {SOURCE}")

    outputs = {
        channel: SOURCE.parent / f"stitched_{channel}_images.pptx"
        for channel in CHANNELS
    }

    # Check all output paths before creating any presentations.
    existing = [path for path in outputs.values() if path.exists()]
    if existing:
        filenames = "\n".join(str(path) for path in existing)
        raise FileExistsError(
            "These output files already exist:\n"
            f"{filenames}\n"
            "Rename them or change the output names before running again."
        )

    matches = {channel: [] for channel in CHANNELS}

    # Search the source folder once for all three image types.
    for path in SOURCE.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".tif", ".tiff"}:
            continue

        for channel in CHANNELS:
            reason = match_reason(path, channel)
            if reason is not None:
                matches[channel].append((path, reason))

    for channel in CHANNELS:
        channel_matches = sorted(
            matches[channel],
            key=lambda item: natural_sort_key(item[0]),
        )

        if not channel_matches:
            print(f"\nNo stitched-{channel} TIFF files found; skipping.")
            continue

        print(
            f"\nFound {len(channel_matches)} "
            f"stitched-{channel} TIFF files:\n"
        )
        for path, reason in channel_matches:
            print(f"{path.relative_to(SOURCE)} [{reason}]")

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        for number, (path, _) in enumerate(channel_matches, start=1):
            print(
                f"\n[{channel}] Adding slide "
                f"{number}/{len(channel_matches)}: {path.name}"
            )
            try:
                add_slide(presentation, path)
            except Exception as error:
                raise RuntimeError(
                    f"Could not create {channel} slide for: {path}"
                ) from error

        presentation.save(outputs[channel])
        print(
            f"\nSaved {len(channel_matches)} slides to:\n"
            f"{outputs[channel]}"
        )

if __name__ == "__main__":
    main()